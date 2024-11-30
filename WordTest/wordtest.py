from PyPDF2 import PdfReader
from docx import Document
import random
from docx2pdf import convert



def extract_lines_after_keyword(text, keyword, num_lines=30):
    """
    멀티라인 텍스트에서 특정 키워드 이후의 N 라인을 추출합니다.

    Args:
        text (str): 멀티라인 텍스트.
        keyword (str): 검색할 키워드.
        num_lines (int): 키워드 이후에 가져올 라인의 수 (기본값: 30).

    Returns:
        list: 키워드 이후의 N 라인을 포함한 리스트.
    """
    # 텍스트를 줄 단위로 분리
    lines = text.splitlines()

    # 키워드가 있는 줄 찾기
    for i, line in enumerate(lines):
        if keyword in line:
            # 키워드 이후의 라인 추출
            return lines[i + 1:i + 1 + num_lines]

    # 키워드가 없으면 빈 리스트 반환
    return []




def extract_second_word(text):
    """
    문자열에서 두 번째 단어를 추출합니다.

    Args:
        text (str): 입력 문자열.

    Returns:
        str: 두 번째 단어. 두 번째 단어가 없으면 빈 문자열 반환.
    """
    # 문자열을 공백 기준으로 분리
    words = text.split()
    
    # 두 번째 단어가 있으면 반환, 없으면 빈 문자열 반환
    return words[1] if len(words) > 1 else ""




def validate_input(user_input):
    try:
        # 입력을 공백으로 분리하여 숫자 리스트 생성
        numbers = list(map(int, user_input.split()))
        
        # 숫자가 1개 또는 2개인지 확인
        if len(numbers) not in [1, 2]:
            raise ValueError("숫자는 1개 또는 2개만 입력해야 합니다.")

        # 숫자가 1~60 범위에 있는지 확인
        for num in numbers:
            if not (1 <= num <= 60):
                raise ValueError(f"{num}는 1~60 사이의 숫자가 아닙니다.")

        return numbers
    except ValueError as e:
        # 오류 메시지 출력
        print(f"입력 오류: {e}")






# 프로그램 실행
if __name__ == "__main__":
    user_input = input("생성할 Day의 숫자를 1개 또는 2개(숫자 사이는 띄어쓰기로) 입력하세요 (예: 5 또는 10 20): ")
    days = validate_input(user_input)    

    # 책 내용 읽기
    reader = PdfReader("C:/pythonprojects/WordTest/뜯어먹는_수능_1등급_기본_영단어_1800_개정판__단어_리스트_241130_083500.pdf")
    # 작성할 시험지 양식 불러오기
    doc = Document('C:/pythonprojects/WordTest/Word Test Format.docx')

    booktext = ""
    for page in reader.pages:
        booktext += page.extract_text()

    result = []
    for day in days:
        keyword = f"DAY {day}"
        dayRes = extract_lines_after_keyword(booktext, keyword, num_lines=30)
        result += dayRes

    wordlist = []
    for txt in result:
        wordlist.append(extract_second_word(txt))

    random.shuffle(wordlist)

    cRow=1
    bCol = True
    half = int(len(wordlist)/2)
    half

    table = doc.tables[0]

    cRow = 1
    for word in wordlist[:half]:
        table.cell(cRow, 1).text = word
        cRow += 1

    cRow = 1
    for word in wordlist[half:]:
        table.cell(cRow, 4).text = word    
        cRow += 1

    strdays = list(map(str, days))
    filename = f'C:/Users/huije/Desktop/단어시험 Day{'_'.join(strdays)}'

    doc.save(f'{filename}.docx')

    inputFile = f'{filename}.docx'
    outputFile = f'{filename}.pdf'

    convert(inputFile, outputFile)